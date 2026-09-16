# pip install python-pptx
from pptx import Presentation
from pptx.util import Pt

slides = [
    ("Leak Detection for Liquid-Cooled PCBs\nin Datacenter Systems",
     ["Best practices: detection reliability, cost, board integration",
      "Igor Gurevich — Sept 2026"]),
    ("Why Leak Detection Matters",
     ["Direct-to-chip liquid cooling is mainstream for high-TDP boards",
      "One undetected leak = board loss, rack downtime, cascading damage",
      "Goals: fast detection, low false positives, low BOM cost, BMC-actionable"]),
    ("Method 1: Rope/Cable Sensors",
     ["Cable routed under cold plates, tubing, chassis low points",
      "Detects liquid anywhere along length (R or C change)",
      "+ Wide coverage, flexible, reusable",
      "– Condensation false positives, routing complexity, moderate cost",
      "Best for: system/chassis-level protection"]),
    ("Method 2: Interdigitated PCB Traces",
     ["Exposed comb/grid copper at high-risk locations; liquid bridges gap",
      "+ Near-zero BOM cost, standard PCB process, pinpoint, fast",
      "– Limited coverage, flux contamination risk, on-board only",
      "Best for: cold plate ports, QD fittings, under connectors"]),
    ("Method 3: Point (Spot) Sensors",
     ["Packaged discrete sensors at key locations",
      "+ High sensitivity, simple readout, off-the-shelf (Honeywell, Sensata)",
      "– Tiny coverage area; many units needed for distributed protection",
      "Best for: critical single-point risks"]),
    ("Method 4: Capacitive Sensors",
     ["Detect liquid via capacitance change on pads",
      "+ Rejects condensation, works through conformal coating",
      "– More complex electronics, higher cost",
      "Best for: environments intolerant of false positives"]),
    ("Comparison Matrix",
     ["Rope/Cable:  High reliability | Mod cost | Large coverage",
      "Interdigitated: Med-High | Low cost | Localized",
      "Point Sensors: Medium | Low cost | Localized",
      "Capacitive: High | Mod-High cost | Localized/Medium"]),
    ("Best Practices",
     ["Redundancy: wide-area rope + point-of-risk traces",
      "Placement: under ports, along coolant lines, gravity low points",
      "Validate sensors against coolant chemistry (PG25 vs dielectric)",
      "BMC integration: debounce logic → alert, pump stop, valve isolation"]),
    ("Recommended Architecture",
     ["PCB comb traces at cold plate fittings & QDs (~free in BOM)",
      "Rope sensor along manifold/tubing path & chassis low points",
      "Signals → BMC with condensation-rejection logic",
      "Confirmed leak → automated flow isolation + alert",
      "Industry practice: Dell/HPE/Lenovo hybrid approach"]),
    ("References",
     ["Schneider Electric — Liquid Leak Detection in Datacenters",
      "Honeywell — Liquid Leak Detection Sensors",
      "ASUS — Liquid Cooling Solutions for Servers whitepaper"]),
]

prs = Presentation()
layout = prs.slide_layouts[1]  # Title and Content
for title, bullets in slides:
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    tf.text = bullets[0]
    for b in bullets[1:]:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(18)

prs.save("leak_detection_liquid_cooled_pcb.pptx")
print("Saved leak_detection_liquid_cooled_pcb.pptx")
